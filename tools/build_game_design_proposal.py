from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "exports" / "歷史文物守護行動_遊戲流程與機制設計企劃書.docx"
TMP = ROOT / "exports" / "game-design-doc-assets"
TMP.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
BLUE = "DCECF4"
PALE = "F4F8FA"
GOLD = "D29B36"
GRAY = "D9D9D9"
TEXT = "243746"
WHITE = "FFFFFF"
FONT = "Microsoft JhengHei"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(row):
    set_repeat_table_header(row)
    prevent_row_split(row)


def set_table_borders(table, color=GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_run_font(run, name=FONT, size=None, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 頁")
    set_run_font(tail, size=9, color="666666")


def add_paragraph(doc, text="", style=None, bold_lead=None, align=None, space_after=5):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_table(doc, headers, rows, widths=None, font_size=8.7, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 100, 100, 100, 100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2:
            for cell in row.cells:
                set_cell_shading(cell, PALE)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 and i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color="666666")
    return p


def set_image_alt(shape, description):
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", description)


def add_heading(doc, text, level=1, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, color="000000")
    p.paragraph_format.keep_with_next = True
    return p


def convert_webp(source, destination, max_width=1600):
    image = Image.open(source).convert("RGB")
    if image.width > max_width:
        ratio = max_width / image.width
        image = image.resize((max_width, int(image.height * ratio)), Image.Resampling.LANCZOS)
    image.save(destination, quality=94)


def make_flowchart(path):
    width, height = 1800, 760
    img = Image.new("RGB", (width, height), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msjh.ttc"
    bold_path = r"C:\Windows\Fonts\msjhbd.ttc"
    title_font = ImageFont.truetype(bold_path, 42)
    box_font = ImageFont.truetype(bold_path, 25)
    small_font = ImageFont.truetype(font_path, 19)
    draw.text((60, 30), "故事模式完整流程", font=title_font, fill="#17324D")
    steps = [
        ("前置劇情", "了解文物失竊"), ("玩家取名", "建立個人代入"), ("行動地圖", "選擇已開放關卡"), ("選擇道具", "一般一件  最終關兩件"),
        ("課前預測", "記錄原本想法"), ("探索兩處", "完成資料判讀"), ("隨機事件", "整局三至五次"), ("證據戰鬥", "判斷再選依據"),
        ("區域學習單", "用證據修正想法"), ("取回文物", "獲得知識卡"), ("開放下一關", "守護心持續保留"), ("最終結局", "回饋與守護誓詞"),
    ]
    box_w, box_h = 380, 145
    xs = [60, 500, 940, 1380]
    ys = [115, 330, 545]
    for idx, (title, subtitle) in enumerate(steps):
        row, col = divmod(idx, 4)
        x, y = xs[col], ys[row]
        fill = "#E8F2F7" if row != 2 else "#FFF4D6"
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=22, fill=fill, outline="#2B5873", width=4)
        draw.text((x + 24, y + 22), f"{idx + 1:02d}  {title}", font=box_font, fill="#17324D")
        draw.text((x + 24, y + 78), subtitle, font=small_font, fill="#40576A")
        if col < 3:
            ax = x + box_w + 10
            ay = y + box_h // 2
            draw.line((ax, ay, ax + 42, ay), fill="#D29B36", width=7)
            draw.polygon([(ax + 42, ay), (ax + 26, ay - 11), (ax + 26, ay + 11)], fill="#D29B36")
        elif row < 2:
            ax = x + box_w // 2
            ay = y + box_h + 8
            draw.line((ax, ay, ax, ay + 42), fill="#D29B36", width=7)
            draw.polygon([(ax, ay + 42), (ax - 11, ay + 26), (ax + 11, ay + 26)], fill="#D29B36")
    img.save(path)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.75)
section.right_margin = Cm(1.75)

styles = doc.styles
styles["Normal"].font.name = FONT
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
styles["Normal"].paragraph_format.line_spacing = 1.35
styles["Normal"].paragraph_format.space_after = Pt(5)
for name, size in (("Title", 25), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
    style = styles[name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Title" or True
    style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = FONT
    styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles[style_name].font.size = Pt(10.5)

footer = section.footer.paragraphs[0]
add_page_number(footer)

hero_png = TMP / "cover-hero.png"
map_png = TMP / "adventure-map.png"
flow_png = TMP / "game-flow.png"
convert_webp(ROOT / "assets" / "guardian-archive-hero.webp", hero_png)
convert_webp(ROOT / "assets" / "adventure-map-taiwan.webp", map_png)
make_flowchart(flow_png)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
r = p.add_run("歷史文物守護行動")
set_run_font(r, size=28, bold=True, color="000000")
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("遊戲流程與機制設計企劃書")
set_run_font(r, size=20, bold=True, color="000000")
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p3.add_run("國小五年級社會科網頁探究遊戲")
set_run_font(r, size=12, color="444444")
cover_shape = doc.add_picture(str(hero_png), width=Cm(14.8))
set_image_alt(cover_shape, "玩家守護員與文物嚮導檔檔站在博物館典藏空間")
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption(doc, "目前遊戲主視覺  玩家與文物嚮導檔檔共同展開守護行動")
add_table(doc, ["文件項目", "內容"], [
    ["企劃範圍", "目前可執行的十關故事模式 十單元教師模式 題庫與戰鬥系統"],
    ["主要玩家", "國小五年級學生 亦適合五至六年級使用"],
    ["遊玩平台", "桌面與行動裝置瀏覽器"],
    ["文件版本", "2026年10月14日 現行程式版本"],
], widths=[3.4, 13.8], font_size=9.3)

doc.add_page_break()
add_heading(doc, "企劃摘要", 1)
add_paragraph(doc, "歷史文物守護行動是一款把國小社會科探究歷程轉化為冒險任務的網頁遊戲。十個故事關卡依序對應五上、五下十個學習單元，玩家從臺灣的位置與生活樣貌出發，走過物產、歷史、能源、經濟與永續家園。每一關都要求玩家先做預測，再操作線索、完成判斷、選出支持判斷的知識，最後以區域學習單整理學習結果。")
add_paragraph(doc, "本作的核心設計不是讓學生只靠記憶選答案，而是把課堂常用的提問、蒐證、比較、判斷、說理與反思安排成固定遊戲循環。戰鬥提供動機與立即回饋；課前預測、探索驗證、成立條件及區域學習單則保留可被教師觀察的學習行為。")
add_heading(doc, "目前完成範圍", 2)
add_bullets(doc, [
    "十關連續故事與完整結局，每關一對一對應學習單元，並包含兩個探索點、一場戰鬥、一件文物及一張知識卡。",
    "五上與五下共十個單元，每單元十五題，共一百五十題；每場依敵人血量抽題且不重複。",
    "戰鬥錯誤選項會引用其他單元的知識，並以不同回饋說明為何不能回答目前問題。",
    "五種可選道具、十五種隨機事件、逐關提高的血量與攻擊力，以及最終關雙道具配置。",
    "教師課堂模式可直接選擇十個單元、設定十至三十分鐘、暫停講解並匯出本機學習摘要。",
    "同機雙人搶答模式，與故事進度及單人存檔分開。",
])
add_heading(doc, "文件閱讀方式", 2)
add_table(doc, ["章節", "用途"], [
    ["一至四", "掌握產品定位 世界觀及完整遊戲流程"],
    ["五至十二", "查閱關卡 戰鬥 題庫 道具 事件與平衡規格"],
    ["十三至十六", "了解課堂使用 雙人模式 介面與技術架構"],
    ["十七至十九", "執行測試 教學評估及後續迭代"],
], widths=[3.2, 14], font_size=9.4)

add_heading(doc, "一 專案定位與設計目標", 1, page_break=True)
add_heading(doc, "產品定位", 2)
add_paragraph(doc, "本作定位為可在課堂中使用，也能讓學生自行完成的社會科探究冒險遊戲。故事模式提供連續情節與成長壓力；教師模式則讓教師依課程進度直接選擇單元，兩種入口共用同一套題庫、探索任務、戰鬥回饋與區域學習單。")
add_table(doc, ["面向", "設計規格"], [
    ["核心受眾", "國小五年級學生 文字與選項以五年級可理解程度撰寫"],
    ["次要受眾", "六年級複習 教師課堂活動 小組競賽"],
    ["平台", "免安裝的靜態網頁 支援滑鼠 鍵盤與觸控"],
    ["內容範圍", "國小社會五上與五下十個單元的探究與題庫內容"],
    ["主要情境", "課前導入 課中形成性評量 課後複習 自主遊玩"],
    ["單局結構", "預測 探索 隨機事件 戰鬥 區域學習單 獎勵"],
], widths=[3.4, 13.8], font_size=9.4)
add_heading(doc, "設計目標", 2)
add_numbered(doc, [
    "讓學生看懂題目後做出判斷，而不是記住正確答案的位置。",
    "讓學生分辨答案與支持答案的知識依據，建立主張與證據的連結。",
    "讓答錯成為可修正的學習事件，提供針對選項的提示、傷害回饋與重新作答機會。",
    "讓教師能在十至三十分鐘內取得可討論的預測、錯誤類型及區域學習單。",
    "以逐關成長、文物蒐集與劇情推進維持學生完成整套內容的動機。",
])
add_heading(doc, "核心設計原則", 2)
add_table(doc, ["原則", "落實方式"], [
    ["先思考再作答", "先完成預測與探索 才能進入戰鬥"],
    ["答案必須有依據", "答對四選一後 還要選出成立條件"],
    ["錯誤可以修正", "錯誤選項有個別說明 並重新排列全部選項"],
    ["難度逐步成長", "敵人血量 攻擊力 守護心上限及道具效果逐關調整"],
    ["課堂可觀察", "記錄預測 錯誤概念 證據作答與區域學習單"],
    ["尊重多元觀點", "劇情反對刪除不同說法 題目要求比較資料來源與影響者"],
], widths=[4, 13.2], font_size=9.2)

add_heading(doc, "二 教學問題與遊戲解法", 1, page_break=True)
add_paragraph(doc, "傳統選擇題容易只留下答對或答錯的結果，教師不一定能看見學生為何選擇，也難以確認學生能否把課本知識用在新情境。本作把一題拆成判斷與成立條件兩步，並在每關前後加入預測與區域學習單，使作答結果能回到課堂討論。")
add_table(doc, ["課堂需求", "遊戲機制", "教師可觀察內容"], [
    ["了解學生原有想法", "課前預測 同意 不確定 不同意", "學生在接觸資料前的初始立場"],
    ["練習讀圖與比較", "兩個探索點與資料判讀任務", "學生能否找出和問題直接相關的線索"],
    ["說明判斷理由", "答對後再選成立條件", "學生能否區分答案與支持依據"],
    ["辨認錯誤概念", "錯誤選項混入其他單元知識並提供不同回饋", "學生能否辨認知識適用的問題與情境"],
    ["促進反思與遷移", "區域學習單與確定程度", "學生能否用證據修正想法並連結生活"],
    ["保留形成性資料", "本機摘要 列印與 JSON 匯出", "預測 探索 證據 錯誤與區域學習單紀錄"],
], widths=[4.2, 6.5, 6.5], font_size=8.7)

add_heading(doc, "三 世界觀與完整劇情", 1, page_break=True)
add_heading(doc, "故事前提", 2)
add_paragraph(doc, "博物館閉館後，十件核心文物遭抹除者打散。它們分別連結五上、五下十個社會科單元。文物嚮導檔檔請玩家成為守護員，沿著臺灣形狀的行動地圖追查線索，找回文物並重建從地理、生活、歷史到永續家園的臺灣故事。")
add_heading(doc, "主要角色", 2)
add_table(doc, ["角色", "功能", "學習意義"], [
    ["玩家守護員", "取名後成為故事主角 完成調查與戰鬥", "讓學習行動以自己的名字被回饋"],
    ["檔檔", "說明任務 提供提示 支援防禦與課堂導引", "代表查證 資料整理與同儕協作"],
    ["抹除者", "刪除不同資料 讓線索只剩單一說法", "具象化過度簡化 忽略來源與排除不同觀點的問題"],
], widths=[3, 7, 7.2], font_size=9.1)
add_heading(doc, "四段式故事弧線", 2)
add_table(doc, ["段落", "關卡", "劇情任務", "概念推進"], [
    ["建立島嶼視角", "第一至三關", "修復位置 生活與物產資料", "位置交流 多元生活 資源永續"],
    ["重建歷史證據", "第四至七關", "整理大航海 移民與開港變化", "比較時序 來源 因果與多方影響"],
    ["連結現代生活", "第八至九關", "修復電力圖與生活選擇帳簿", "能源風險 消費責任與循環利用"],
    ["完成守護行動", "第十關", "面對抹除者 形成永續家園方案", "多元觀點 共識 補救與持續檢討"],
], widths=[3.1, 2.8, 6.2, 5.1], font_size=8.9)
add_paragraph(doc, "最終戰後，抹除者理解不同線索可以拼成更完整的故事，並成為故事記錄員。結局以守護員誓詞收束，讓查清資料、尊重不同想法及保護自然與文物成為遊戲外仍可執行的行動。")

add_heading(doc, "四 整體遊戲流程", 1, page_break=True)
flow_shape = doc.add_picture(str(flow_png), width=Cm(17.2))
set_image_alt(flow_shape, "故事模式十二個主要步驟流程圖")
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption(doc, "圖一  故事模式從前置劇情到最終回饋的完整流程")
add_paragraph(doc, "故事模式採線性解鎖。玩家可重玩已開放關卡，但尚未完成前一關時不能提前進入後續關卡。每次進關會增加進入次數種子，重新排列探索點、線索選項、抽出的題目及答案位置。")
add_heading(doc, "成功與失敗路徑", 2)
add_table(doc, ["狀態", "系統處理", "玩家下一步"], [
    ["完成判斷與成立條件", "對敵人造成基礎一點傷害 道具或事件可加成", "進入下一題或區域學習單"],
    ["判斷答錯", "立即承受敵人攻擊 顯示該選項說明 全部選項重新排列", "重新閱讀同題並作答"],
    ["成立條件答錯", "顯示該知識與本題的差異 不扣守護心", "保留原判斷並重選依據"],
    ["敵人血量歸零", "播放勝利與文物回收演出", "完成區域學習單並領取知識卡"],
    ["守護心歸零", "播放戰敗演出後回到地圖", "再次進關時恢復目前守護心上限"],
], widths=[3.6, 8.3, 5.4], font_size=8.9)

add_heading(doc, "五 單關流程規格", 1, page_break=True)
add_table(doc, ["階段", "玩家行動", "系統規則", "教學目的"], [
    ["一 關卡簡報", "閱讀本關衝突與學習資訊", "顯示課程位置 目標 操作 難度與剩餘守護心", "建立任務脈絡"],
    ["二 道具配置", "一般關選一件 最終關選一至兩件", "只顯示已解鎖道具 每場各可用一次", "規畫風險與資源"],
    ["三 課前預測", "選同意 不確定或不同意", "無標準答案並記入學習紀錄", "喚起先備想法"],
    ["四 探索", "依隨機順序調查兩個場景", "每點完成三選一驗證 正確才記錄線索", "讀圖 分類 排序與比較"],
    ["五 隨機事件", "閱讀事件並選擇策略或回答小題", "整局三至五場戰鬥前觸發 同場不重複", "應用知識與資源判斷"],
    ["六 證據戰鬥", "先選答案 再選成立條件", "答錯受傷並重排 答對才能攻擊", "連結主張與證據"],
    ["七 區域學習單", "重看課前問題 寫至少八字並選確定程度", "以課前問題為主標題 作答提示為較小副標題", "修正想法並形成說明"],
    ["八 關卡獎勵", "閱讀文物說明與知識卡", "標記完成 解鎖下一關 保存進度", "統整內容並推進故事"],
], widths=[2.7, 4.4, 6.3, 3.8], font_size=8.5)

add_heading(doc, "六 十關內容與難度規格", 1, page_break=True)
map_shape = doc.add_picture(str(map_png), width=Cm(17.2))
set_image_alt(map_shape, "臺灣形狀的十關遊戲行動地圖")
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption(doc, "圖二  十關行動地圖與文物修復路線")
levels = [
    [1, "臺灣！我來了", "五上第一單元", "島嶼地圖桌 考古證據箱", "航線迷霧 3HP 攻1 心4", "臺灣探索圖 島嶼觀察羅盤"],
    [2, "島嶼上的生活樣貌", "五上第二單元", "文化比較牆 共同生活站", "單一印象墨團 4HP 攻1 心4", "多元生活圖冊 文化聆聽石"],
    [3, "臺灣的物產資源", "五上第三單元", "山林農場 海岸漁港", "耗竭砂暴 4HP 攻1 心5", "物產共生圖 永續量尺"],
    [4, "看見福爾摩沙", "五上第四單元", "世界航線室 經營紀錄牆", "時序海怪 5HP 攻2 心5", "海權時序圖 證據印章"],
    [5, "發現臺灣之美", "五上第五單元", "問題設計桌 資料檢核站", "片面結論霧 5HP 攻2 心6", "臺灣探究手冊 行動筆記"],
    [6, "唐山到臺灣", "五下第一單元", "移民航線圖 水圳聚落桌", "渡海亂流 6HP 攻2 心6", "渡臺家族簿 水圳協作尺"],
    [7, "開港通商後的臺灣", "五下第二單元", "港市貨物臺 建設時序站", "開港煙霧 7HP 攻3 心7", "開港貿易圖 因果時序鏡"],
    [8, "臺灣電力的發展", "五下第三單元", "發電條件牆 風險分配桌", "能源失衡獸 8HP 攻3 心7", "臺灣電力圖 能源評估表"],
    [9, "經濟與生活", "五下第四單元", "消費查證站 循環共享桌", "消費誘惑影 8HP 攻3 心8", "生活選擇帳簿 風險查證章"],
    [10, "打造永續的家園", "五下第五單元", "多元方案牆 共識行動桌", "抹除者 9HP 攻4 心8", "永續家園圖卷 守護者徽章"],
]
add_table(doc, ["關", "名稱", "課程主題", "探索點", "敵人數值", "文物與知識卡"], levels, widths=[1, 3.3, 3.2, 3.9, 3.3, 3.8], font_size=7.8)
add_heading(doc, "各關設計重點", 2)
level_details = [
    ("第一關 臺灣！我來了", "運用地圖、地形與考古證據，理解臺灣的位置、島嶼形成及史前生活。"),
    ("第二關 島嶼上的生活樣貌", "比較族群文化、社會規範與法律，練習從多方觀點處理共同生活的衝突。"),
    ("第三關 臺灣的物產資源", "從經濟、生活與環境三面向評估物產利用，思考資源如何長久使用。"),
    ("第四關 看見福爾摩沙", "依時間、地點與資料來源重建大航海時代，避免只保留單一立場。"),
    ("第五關 發現臺灣之美", "提出可查證問題、檢查資料來源與限制，並設計可持續修正的行動。"),
    ("第六關 唐山到臺灣", "追查移民、開墾、水利及族群互動，說明改變對不同人群造成的影響。"),
    ("第七關 開港通商後的臺灣", "連結開港、經濟變化、戰爭與建設，評估不同群體的受益與代價。"),
    ("第八關 臺灣電力的發展", "比較各種發電方式的條件、效益、風險與受影響者。"),
    ("第九關 經濟與生活", "在消費、理財、共享及循環利用情境中查證風險並負責任地選擇。"),
    ("第十關 打造永續的家園", "整合前九關證據，形成含共同目標、補救及檢查方式的永續方案；最終戰可帶兩件道具。"),
]
for title, body in level_details:
    add_heading(doc, title, 3)
    add_paragraph(doc, body)

add_heading(doc, "七 探索與線索系統", 1, page_break=True)
add_paragraph(doc, "每關固定設置兩個探索點，但每次進關時會重新排列出現順序。玩家點選場景熱點後進入三選一驗證，選項也會重新排列，因此不能靠位置記憶。答對後，該探索點的完整線索才會寫入線索筆記。兩條線索都完成後，戰鬥入口才會開放。")
add_table(doc, ["規則", "目前實作"], [
    ["探索數量", "每關兩處 共二十個故事探索點"],
    ["任務操作", "讀圖 比較 分類 排序 判斷或多觀點分析"],
    ["選項設計", "每題三個選項 正確位置依進關次數重新排列"],
    ["錯誤處理", "顯示線索和問題關係不大的提示 可立即重選"],
    ["完成條件", "兩處都答對並寫入線索筆記"],
    ["教學功能", "要求學生先接觸資料 再進入抽象題目與戰鬥"],
], widths=[4, 13.2], font_size=9.2)

add_heading(doc, "八 證據戰鬥系統", 1, page_break=True)
add_heading(doc, "基本回合", 2)
add_numbered(doc, [
    "系統從本單元十五題中依敵人血量抽出題目，同一場不重複。",
    "題目與四個選項出現，正確答案位置重新排列。",
    "玩家選錯時立即受到敵人攻擊，畫面顯示該錯誤選項的專屬說明。",
    "同一題的四個選項全部保留並重新排列，玩家重新判斷。",
    "玩家選對後進入第二步，從三項知識中選出支持目前判斷的成立條件。",
    "成立條件答錯不扣守護心，系統解釋該知識屬於哪個主題，玩家可重選。",
    "判斷與成立條件都正確後，玩家攻擊敵人並得到完整解析。",
])
add_heading(doc, "傷害與結束條件", 2)
add_table(doc, ["項目", "公式或規則"], [
    ["玩家基礎傷害", "每次完成判斷與成立條件造成一點傷害"],
    ["敵人傷害", "依關卡攻擊力扣除守護心 護盾或檔檔支援可抵銷"],
    ["額外傷害", "證據印章或事件可提高下一次正確作答傷害"],
    ["敵人戰敗", "敵人血量歸零後播放最後一擊與文物回收演出"],
    ["玩家戰敗", "守護心歸零後播放戰敗演出並返回地圖"],
    ["題目耗盡", "尚未擊敗敵人時視為戰敗 避免無限作答"],
], widths=[4, 13.2], font_size=9.2)
add_heading(doc, "回饋層級", 2)
add_table(doc, ["層級", "回饋內容"], [
    ["選項回饋", "每個答案分別說明為何成立或哪個條件被忽略"],
    ["迷思紀錄", "記錄位置錯置 過度概括 單因判斷 證據不足等類型"],
    ["成立條件回饋", "指出知識能否支持本題 且文字不得與答案選項相同"],
    ["完整解析", "說明題目線索 因果關係 可推到的範圍與限制"],
    ["視聽回饋", "答對 攻擊 受傷 防禦 治療 勝敗與最後一擊皆有動畫和音效"],
], widths=[4, 13.2], font_size=9.1)

add_heading(doc, "九 題庫與內容規格", 1, page_break=True)
add_paragraph(doc, "題庫涵蓋五上及五下十個單元，每單元十五題。每五個核心案例各衍生情境判斷、證據判讀與迷思修正三種題型，形成一百五十題。故事模式依關卡指定題庫；教師模式可直接選擇任何單元。")
add_table(doc, ["學期", "單元", "探究重點"], [
    ["五上", "臺灣 我來了", "位置 地形 考古證據"],
    ["五上", "島嶼上的生活樣貌", "族群文化 規範 法律 權利"],
    ["五上", "臺灣的物產資源", "農林漁牧礦 永續與地方轉型"],
    ["五上", "看見福爾摩沙", "航路 時序 多元史料與大航海影響"],
    ["五上", "發現臺灣之美", "提問 蒐證 分析限制與行動"],
    ["五下", "唐山到臺灣", "移民 水利 港市與族群互動"],
    ["五下", "開港通商後的臺灣", "開港 建省 戰爭與現代化影響"],
    ["五下", "臺灣電力的發展", "能源條件 風險 代價與環境公平"],
    ["五下", "經濟與生活", "消費 理財 共享與循環經濟"],
    ["五下", "打造永續的家園", "多元觀點 共識 能源方案與行動評估"],
], widths=[2.5, 5.4, 9.3], font_size=8.8)
add_heading(doc, "題目品質規則", 2)
add_bullets(doc, [
    "題幹最多五十字，四個選項各自最多二十字，選項長度差距不超過八字。",
    "同題四個選項必須不同，並檢查過度相似的文字組合。",
    "每題至少一個錯誤選項取自其他單元，證據判讀與迷思修正題通常使用三個跨單元干擾項。",
    "避免不必要的進階用語，但保留板塊、季風、永續等課程必要名詞。",
    "每個選項有專屬回饋及迷思標記，不能只以答對或答錯結束。",
    "成立條件共有三項，一項正確，兩項干擾知識取自其他單元。",
    "任何成立條件不得與本題四個答案選項完全相同。",
    "每場抽題數等於敵人血量，使用隨機索引且不得重複。",
])

add_heading(doc, "十 道具與配置系統", 1, page_break=True)
add_table(doc, ["道具", "解鎖", "前期效果", "強化效果", "定位"], [
    ["航海羅盤", "第一關", "第一至四關排除一個錯誤選項", "第五關起排除兩個", "降低選項壓力"],
    ["線索放大鏡", "第二關", "顯示本題提示", "第五關起提示並排除一個錯誤選項", "支援理解"],
    ["修復工具箱", "第一關", "第一至三關回復一顆守護心", "第四至六關回復兩顆 第七關起三顆", "續戰與容錯"],
    ["時序沙漏", "第五關", "抵擋一次敵人攻擊", "第七關起抵擋兩次", "高攻擊關卡防護"],
    ["證據印章", "第七關", "下一次答對造成兩點傷害", "最終關造成三點傷害", "縮短高血量戰鬥"],
], widths=[3, 2, 4.4, 5, 2.8], font_size=8.4)
add_paragraph(doc, "一般關卡必須選一件道具，最終頭目最多可攜帶兩件。每件道具每場只能使用一次；若工具箱在守護心全滿時使用，系統不會消耗次數。道具效果隨關卡強化，使玩家能應付後期提高的敵人攻擊。")

add_heading(doc, "十一 隨機事件系統", 1, page_break=True)
add_paragraph(doc, "新遊戲會先規畫三至五個事件，隨機放在不同關卡的戰鬥前。同一場戰鬥不會重複觸發，整局也不會抽到同一個事件兩次。事件類別權重為幸運與保護百分之三十、知識事件百分之二十五、策略選擇百分之二十五、小危機百分之二十。")
events = [
    ["文物修復站", "幸運與保護", "回復一顆守護心"], ["檔案館密函", "知識事件", "第一題排除一個錯誤選項"],
    ["守護隊補給箱", "幸運與保護", "借用額外道具 無可借道具時回復守護心"], ["時空順風", "幸運與保護", "第一次答對額外一點傷害"],
    ["迷霧突襲", "小危機", "失去一顆守護心 第一題顯示提示"], ["岔路上的腳印", "策略選擇", "選回復守護心或削弱敵人"],
    ["真假線索挑戰", "知識事件", "答對獲得一次護盾 答錯不扣心"], ["旅行商人的交換", "策略選擇", "選修復包或證據徽章"],
    ["守護員來援", "幸運與保護", "首次歸零時回復一顆並繼續"], ["文物的回音", "知識事件", "已修復文物削弱敵人 無文物時給護盾"],
    ["缺字的文物標籤", "知識事件", "答對使敵人開場少一點血量"], ["博物館突然停電", "小危機", "選第一題排除兩項或獲得護盾"],
    ["檔檔的分工計畫", "策略選擇", "選第一題提示或降低首次傷害"], ["迷霧交換所", "小危機", "可用一顆守護心換敵人兩點血量"],
    ["時空車站幸運票", "幸運與保護", "隨機獲得治療 護盾或額外傷害"],
]
add_table(doc, ["事件", "類別", "主要效果"], events, widths=[5, 4.2, 8], font_size=8.7)

add_heading(doc, "十二 難度 成長與資源平衡", 1, page_break=True)
add_table(doc, ["關卡", "難度", "敵人血量", "敵人攻擊", "守護心上限", "抽題數"], [
    [1, "見習", 3, 1, 4, 3], [2, "初級", 4, 1, 4, 4], [3, "初級＋", 4, 1, 5, 4], [4, "進階", 5, 2, 5, 5],
    [5, "進階＋", 5, 2, 6, 5], [6, "挑戰", 6, 2, 6, 6], [7, "高階", 7, 3, 7, 7], [8, "艱難", 8, 3, 7, 8],
    [9, "專家", 8, 3, 8, 8], [10, "最終守護", 9, 4, 8, 9],
], widths=[2, 3.3, 3, 3, 3.2, 2.7], font_size=9)
add_paragraph(doc, "敵人血量由三逐步提高至九，攻擊力由一提高至四。守護心上限也由四提高至八，但成功過關後不自動補滿，因此玩家需要保留修復工具箱或利用事件回復。若戰敗後守護心為零，重新進關時會恢復目前上限，避免永久無法繼續。")
add_heading(doc, "平衡意圖", 2)
add_bullets(doc, [
    "前三關建立規則，攻擊力維持一，讓學生熟悉答錯即受傷與成立條件流程。",
    "第四關開始提高攻擊，迫使玩家開始考慮治療、防禦和提示道具。",
    "第五關後道具同步升級，避免敵人變強卻只延長作答時間。",
    "第十關提供兩件道具，讓最終戰的九題與四點攻擊仍有策略空間。",
    "隨機事件提供短期補償，但觸發次數有上限，不能完全取代題目理解。",
])

add_heading(doc, "十三 教師課堂模式", 1, page_break=True)
add_paragraph(doc, "教師模式不受故事解鎖限制，可直接開啟五上、五下十個單元。教師先選十、二十或三十分鐘，再依當週課程選擇題庫。畫面保留預測、探索、戰鬥與區域學習單，使遊戲可嵌入一節課，而不必先完成前面故事。")
add_table(doc, ["課堂時間", "建議流程", "教師介入重點"], [
    ["十分鐘", "預測兩分鐘 探索三分鐘 戰鬥四分鐘 區域學習單一分鐘", "選一個錯誤選項快速追問理由"],
    ["二十分鐘", "導入三分鐘 探索六分鐘 戰鬥八分鐘 區域學習單三分鐘", "比較兩組學生使用的證據"],
    ["三十分鐘", "導入五分鐘 探索八分鐘 戰鬥十二分鐘 分享與區域學習單五分鐘", "暫停畫面進行全班討論與修正"],
], widths=[3, 8.5, 5.7], font_size=8.9)
add_heading(doc, "教師功能", 2)
add_bullets(doc, [
    "直接選擇十個課程單元及建議課堂時間。",
    "教師暫停功能讓全班在關鍵題目停下討論。",
    "學習摘要統計課前預測、探索驗證與區域學習單數量。",
    "摘要只保存在本機，不含帳號或班級代碼，可列印或匯出 JSON。",
    "最多保留近期八十筆學習紀錄，避免瀏覽器資料無限成長。",
])
add_heading(doc, "課堂提問範例", 2)
add_table(doc, ["時機", "教師提問"], [
    ["課前預測後", "你選同意或不同意的依據是什麼 目前還缺哪一種資料"],
    ["探索答錯時", "這個線索真的能回答題目嗎 它少了時間 地點或比較對象嗎"],
    ["戰鬥選項後", "正確答案和成立條件有什麼不同 哪一項是結論 哪一項是依據"],
    ["區域學習單前", "你原本的想法有沒有改變 哪一條線索最重要"],
], widths=[4, 13.2], font_size=9.1)

add_heading(doc, "十四 雙人對決模式", 1, page_break=True)
add_paragraph(doc, "雙人對決供兩位學生在同一台電腦上搶答，與故事存檔、關卡守護心及解鎖狀態分開。玩家只能選擇目前故事模式已開放的關卡題庫。")
add_table(doc, ["項目", "規格"], [
    ["操作鍵", "玩家一使用 A W S D  玩家二使用 J I K L 或方向鍵"],
    ["閱讀時間", "選項先隱藏二點七秒 再開放搶答"],
    ["玩家守護心", "每人三顆"],
    ["敵人血量", "固定五點"],
    ["答對", "先答對者造成一點傷害"],
    ["答錯", "扣一顆守護心 該玩家本題鎖定 並替對手排除該錯誤選項"],
    ["勝負", "先比較造成傷害 傷害相同時比較剩餘守護心"],
], widths=[4, 13.2], font_size=9.2)
add_paragraph(doc, "此模式適合課堂複習或小組代表賽。教師應先提醒學生以理解優先，因為答錯後該題不能再答，盲目搶按會直接降低勝率。")

add_heading(doc, "十五 介面 回饋與無障礙", 1, page_break=True)
add_table(doc, ["面向", "目前設計"], [
    ["個人化", "首頁輸入一至十二字名字 劇情 戰鬥與結局皆使用玩家名稱"],
    ["響應式版面", "地圖支援手機橫向捲動 主要畫面隨寬度調整"],
    ["文字可讀性", "提供放大文字 題幹與選項限制長度 使用適合五年級的語句"],
    ["鍵盤操作", "單人戰鬥可使用答案快捷鍵 雙人模式配置兩組按鍵"],
    ["狀態提示", "按鈕有文字 狀態與圖示 不只依靠顏色"],
    ["動態效果", "尊重系統的減少動態設定 並縮短動畫等待"],
    ["聲音", "可切換音效 答對 受傷 防禦 治療及勝敗使用不同聲音"],
    ["讀屏資訊", "使用標題 按鈕 標籤 進度條 live status 與 ARIA 描述"],
], widths=[4, 13.2], font_size=9)

add_heading(doc, "十六 技術架構與資料保存", 1, page_break=True)
add_heading(doc, "前端架構", 2)
add_table(doc, ["檔案或模組", "責任"], [
    ["歷史文物守護行動完整版.html", "畫面樣式 遊戲狀態 關卡 戰鬥 事件 動畫與存檔流程"],
    ["question-bank.js", "十個單元共一百五十題 選項 回饋 迷思 成立條件與遷移情境"],
    ["curriculum-data.js", "十單元課程定位 十關教學目標 預測 探索與區域學習單"],
    ["assets", "地圖 關卡背景 角色 敵人 文物 道具 事件與字型資源"],
    ["Excel題庫與轉換工具", "提供可編輯題庫並驗證後產生瀏覽器可讀資料"],
], widths=[5.4, 11.8], font_size=9)
add_paragraph(doc, "遊戲為純前端靜態網站，不需要帳號、資料庫或伺服器 API。瀏覽器載入 HTML、課程資料、題庫與美術資源後即可執行，適合本機開啟或放置於 Firebase Hosting 等靜態網站服務。")
add_heading(doc, "本機存檔", 2)
add_table(doc, ["資料", "用途"], [
    ["玩家名稱", "劇情與戰鬥個人化"], ["解鎖與完成關卡", "控制地圖進度與重玩範圍"],
    ["守護心", "保存跨關資源狀態"], ["事件計畫與已觸發事件", "維持整局三至五次且不重複"],
    ["試玩回饋", "保存結局頁勾選結果"], ["課堂時間與學習紀錄", "支援教師摘要 列印與匯出"],
], widths=[5, 12.2], font_size=9.1)
add_paragraph(doc, "存檔版本目前為第三版，啟動時會處理舊資料。學習紀錄只留在使用者瀏覽器，清除網站資料或按下清除進度後會消失。")

add_heading(doc, "十七 測試與驗收規格", 1, page_break=True)
add_table(doc, ["測試面向", "驗收條件"], [
    ["題庫完整性", "十個單元各十五題 題幹不重複 每題四個唯一選項"],
    ["可讀性", "題幹與選項符合長度限制 不使用不必要的進階詞彙"],
    ["答案亂數", "各關多次進入後 正確答案能出現在不同位置"],
    ["跨單元干擾", "每題至少一個錯誤選項來自其他單元 且不使用制式干擾句"],
    ["探索亂數", "兩個探索點及三個線索位置會改變 正確位置可涵蓋三格"],
    ["戰鬥抽題", "抽題數等於敵人血量 同一場不重複"],
    ["成立條件", "三項文字唯一 一項正確 干擾項來自其他單元 且不得複製答案"],
    ["事件", "十五種唯一事件 權重總和百分之百 每局三至五次"],
    ["道具", "解鎖限制 使用次數 強化效果與最終關兩件配置正確"],
    ["守護心", "成功過關不補滿 治療上限正確 戰敗可返回地圖"],
    ["課堂流程", "預測 探索 戰鬥 區域學習單與摘要資料可完整走通"],
    ["瀏覽器試玩", "主要畫面可操作 無主控台錯誤 重要狀態有可見回饋"],
], widths=[4.2, 13], font_size=8.8)
add_heading(doc, "發布前檢查清單", 2)
add_bullets(doc, [
    "執行所有 check 題庫 事件 探索 道具 守護心 課堂流程及勝敗測試。",
    "以新存檔完成第一關，確認預測 探索 事件 戰鬥 區域學習單與獎勵。",
    "以低守護心進入中後期關卡，確認工具箱與事件不超過上限。",
    "進入第十關確認可選兩件道具 九題不重複及四點敵人傷害。",
    "以桌面與手機寬度檢查地圖、題目、成立條件和區域學習單排版。",
    "確認題庫 Excel 轉換後仍能通過相同驗證，避免覆蓋目前修正。",
])

add_heading(doc, "十八 教學成效評估建議", 1, page_break=True)
add_paragraph(doc, "目前程式已能記錄作答與區域學習單，但是否能改善課堂學習，仍需要透過學生試玩與教師觀察驗證。建議以一個單元進行小規模試教，先檢查學生是否看懂流程，再比較其證據說明品質。")
add_table(doc, ["評估問題", "可蒐集資料", "判讀方式"], [
    ["學生是否理解遊戲流程", "第一次完成預測 探索 戰鬥與區域學習單所需時間", "記錄需要教師協助的步驟"],
    ["學生是否只記答案位置", "多次進入後的正確率與重新排列反應", "比較換位前後表現"],
    ["學生能否區分答案與依據", "成立條件錯誤次數與口頭說明", "檢查是否能說出因為與所以的關係"],
    ["錯誤回饋是否有幫助", "答錯後第二次作答率與訪談", "確認學生是否依提示修正而非猜測"],
    ["內容是否適合五年級", "閱讀停頓 生詞提問 題目完成率", "標記需教師解釋的詞彙"],
    ["學生是否願意持續玩", "結局回饋與關卡完成率", "分辨因難度 文字量或操作造成的中止"],
    ["能否連結課堂", "區域學習單及課後討論內容", "依是否引用線索 比較觀點與提出行動分級"],
], widths=[4.7, 6.2, 6.3], font_size=8.5)
add_heading(doc, "建議的區域學習單評分規準", 2)
add_table(doc, ["等級", "表現描述"], [
    ["四 完整", "提出清楚判斷 引用相關線索 說明因果或限制 並能連結新情境"],
    ["三 達成", "提出判斷並引用相關線索 理由大致合理"],
    ["二 發展中", "有判斷但證據模糊 或只重述選項"],
    ["一 需要協助", "未完成判斷 證據與題目無關 或無法說明"],
], widths=[3.2, 14], font_size=9.2)

add_heading(doc, "十九 風險與後續迭代", 1, page_break=True)
add_table(doc, ["風險", "目前控制", "後續建議"], [
    ["題目內容與課程不完全一致", "依教師手冊主題建題並保留單元標記", "由任課教師逐單元審題並記錄版本"],
    ["學生閱讀負荷過高", "限制題幹與選項長度 提供放大文字", "用五年級實測找出仍需拆句的題目"],
    ["後期失敗率過高", "道具強化 事件補償 最終關雙道具", "蒐集各關錯誤率後再調整數值"],
    ["隨機事件干擾課堂時間", "每局三至五次且只在戰前", "教師模式可新增關閉事件選項"],
    ["教師難以快速掌握全班", "本機摘要可列印與匯出", "未來建立班級代碼與匿名彙整 需先處理隱私"],
    ["雙人模式只適合鍵盤", "提供兩組按鍵與方向鍵", "若要行動裝置使用 再設計雙區觸控按鈕"],
    ["本機資料容易消失", "使用 localStorage 並可匯出摘要", "提醒教師課後立即匯出重要紀錄"],
], widths=[4.1, 6.1, 7], font_size=8.5)
add_heading(doc, "建議迭代順序", 2)
add_numbered(doc, [
    "先完成五年級學生可用性測試，修正看不懂的題幹、按鈕與成立條件。",
    "再進行教師審題與課程對照，確認每單元十五題的內容範圍與難度分布。",
    "第三步蒐集各關答錯率、戰敗率與區域學習單品質，再調整敵人和道具數值。",
    "最後才評估班級彙整、更多合作玩法或線上帳號，避免增加尚未證明需要的系統。",
])

add_heading(doc, "結語", 1, page_break=True)
add_paragraph(doc, "目前版本已建立完整的學習循環：學生先表達原有想法，接著操作資料、驗證線索、完成判斷與成立條件，最後用區域學習單說明新理解。學習單以一開始的問題為主標題，將作答要求縮小放在副標題，幫助學生回想探究焦點。十關故事一對一對應五上、五下十個單元；文物獎勵、跨關守護心、道具與事件提供持續遊玩的動機，教師模式也能直接配合相同的十個課程單元。")
add_paragraph(doc, "下一階段的重點不宜繼續增加題目或功能，而應驗證學生是否看懂規則、能否說出答案成立的原因，以及教師能否在課堂時間內使用摘要進行回饋。完成這三項驗證後，再根據實際資料調整難度與教學介入方式。")

doc.core_properties.title = "歷史文物守護行動 遊戲流程與機制設計企劃書"
doc.core_properties.subject = "國小五年級社會科網頁探究遊戲"
doc.core_properties.comments = "依目前可執行遊戲內容整理"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
