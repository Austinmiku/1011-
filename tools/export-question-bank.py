"""Validate the rich Excel question bank and export browser-ready JS."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

HEADERS = [
    "單元代碼", "學期", "單元", "單元名稱", "學習目標代碼", "題號", "題型", "題目",
    "選項 A", "選項 B", "選項 C", "選項 D", "正確選項", "正確答案", "提示", "答案解析",
    "A 回饋", "B 回饋", "C 回饋", "D 回饋", "A 迷思", "B 迷思", "C 迷思", "D 迷思",
    "證據 A", "證據 B", "證據 C", "證據正解", "證據 A 回饋", "證據 B 回饋", "證據 C 回饋",
    "遷移情境", "來源卷別", "來源檔案", "來源題號",
]
UNIT_KEYS = [f"5A-U{i}" for i in range(1, 5)] + [f"5B-U{i}" for i in range(1, 5)]

def export(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, data_only=False, keep_vba=False)
    sheet = workbook["完整題庫"]
    if [sheet.cell(4, col).value for col in range(1, 36)] != HEADERS:
        raise ValueError("題庫欄位與新版 35 欄格式不符")
    bank = {key: [] for key in UNIT_KEYS}
    for row_number in range(5, 125):
        cells = [sheet.cell(row_number, col) for col in range(1, 36)]
        if any(cell.data_type == "f" for cell in cells):
            raise ValueError(f"第 {row_number} 列含公式")
        record = dict(zip(HEADERS, [cell.value for cell in cells]))
        key = record["單元代碼"]
        if key not in bank:
            raise ValueError(f"第 {row_number} 列單元代碼無效：{key}")
        number = len(bank[key]) + 1
        if record["題號"] != number:
            raise ValueError(f"第 {row_number} 列題號應為 {number}")
        if any(value is None or str(value).strip() == "" for value in record.values()):
            raise ValueError(f"第 {row_number} 列有空白欄位")
        answer_letter = record["正確選項"]
        if answer_letter not in "ABCD":
            raise ValueError(f"第 {row_number} 列正確選項必須是 A-D")
        answer_index = "ABCD".index(answer_letter)
        options = [record[f"選項 {letter}"] for letter in "ABCD"]
        if len(set(options)) != 4:
            raise ValueError(f"第 {row_number} 列含重複選項")
        if options[answer_index] != record["正確答案"]:
            raise ValueError(f"第 {row_number} 列正確答案與選項不一致")
        if record["證據正解"] not in "ABC":
            raise ValueError(f"第 {row_number} 列證據正解必須是 A-C")
        evidence_answer = "ABC".index(record["證據正解"])
        evidence_texts = [record[f"證據 {letter}"] for letter in "ABC"]
        if len(set(evidence_texts)) != 3:
            raise ValueError(f"第 {row_number} 列含重複證據選項")
        if any(text in options for text in evidence_texts):
            raise ValueError(f"第 {row_number} 列的成立條件不可與答案選項相同")
        if any("較多人同意" in str(text) or "選項文字較長" in str(text) for text in evidence_texts):
            raise ValueError(f"第 {row_number} 列仍含制式證據干擾項")
        bank[key].append({
            "id": f"{key}-Q{number:02d}", "unitKey": key,
            "objectiveId": record["學習目標代碼"], "questionType": record["題型"],
            "prompt": record["題目"], "options": options, "answerIndex": answer_index,
            "claim": record["正確答案"], "hint": record["提示"], "explanation": record["答案解析"],
            "optionFeedback": [record[f"{letter} 回饋"] for letter in "ABCD"],
            "misconceptionByChoice": [record[f"{letter} 迷思"] for letter in "ABCD"],
            "evidenceOptions": [{"text": record[f"證據 {letter}"], "correct": i == evidence_answer,
                                 "feedback": record[f"證據 {letter} 回饋"]} for i, letter in enumerate("ABC")],
            "transferContext": record["遷移情境"],
            "sourceFile": record["來源檔案"],
            "sourceQuestionNumber": record["來源題號"],
            "sourceSequence": record["題號"],
        })
    if any(len(questions) != 15 for questions in bank.values()):
        raise ValueError("每個單元必須恰好有 15 題")
    if any(len({q["prompt"] for q in questions}) != 15 for questions in bank.values()):
        raise ValueError("同一單元不得有重複題目")
    destination.write_text(
        "// 由 Excel 新版題庫產生；保留各選項回饋、迷思與證據欄位。\n"
        "window.GUARDIAN_QUESTION_BANK = " + json.dumps(bank, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"已驗證並匯出 {sum(map(len, bank.values()))} 題至 {destination}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Excel .xlsx 題庫檔")
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).resolve().parent.parent / "question-bank.js")
    args = parser.parse_args()
    export(args.source, args.output)
